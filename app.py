import os
import re
import sys
import json
import datetime
from io import BytesIO

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

import requests
from flask import Flask, render_template, request, jsonify, send_file
try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
except Exception:  # pragma: no cover
    Document = None

app = Flask(__name__)

# ---------------------------------------------------------------------------
# Configuration (all overridable via environment variables / Render dashboard)
# ---------------------------------------------------------------------------
LLM_API_KEY = os.environ.get("LLM_API_KEY", "").strip()
LLM_BASE_URL = os.environ.get("LLM_BASE_URL", "https://api.groq.com/openai/v1").rstrip("/")
LLM_MODEL = os.environ.get("LLM_MODEL", "openai/gpt-oss-120b")
LLM_TIMEOUT = int(os.environ.get("LLM_TIMEOUT", "60"))

DEFAULT_UNIT_COUNT = 5


# ---------------------------------------------------------------------------
# Prompt construction
# ---------------------------------------------------------------------------
def build_prompt(data):
    unit_count = int(data.get("unit_count") or DEFAULT_UNIT_COUNT)
    unit_numerals = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"][:unit_count]

    schema_units = ",\n    ".join(
        [
            (
                '{"number": "%s", "title": "short descriptive unit title", '
                '"topics": "A single paragraph in this EXACT style -> '
                'Sub-topic heading: sub point, sub point, sub point. '
                'Second sub-topic heading: sub point, sub point, sub point.", '
                '"assignments": ["assignment 1", "assignment 2", "assignment 3", "assignment 4", "assignment 5"], '
                '"unit_tests": ["unit test question 1", "unit test question 2"]}' % n
            )
            for n in unit_numerals
        ]
    )

    schema = (
        "{\n"
        '  "prerequisites": "one short paragraph listing prerequisite subjects/knowledge",\n'
        '  "overview": "one well written paragraph (5-7 sentences) describing what the course covers and why it matters",\n'
        '  "objectives": ["objective 1", "objective 2", "objective 3", "objective 4", "objective 5", "objective 6"],\n'
        '  "outcomes": ["CO1: ...", "CO2: ...", "CO3: ...", "CO4: ...", "CO5: ..."],\n'
        '  "units": [\n    ' + schema_units + "\n  ],\n"
        '  "text_books": ["Author(s), \\"Title\\", Edition, Publisher, Year.", "..."],\n'
        '  "reference_books": ["Author(s), \\"Title\\", Edition, Publisher, Year.", "..."],\n'
        '  "electronic_resources": ["NPTEL/Coursera/official-docs style resource with a short description", "..."]\n'
        "}"
    )

    example_unit_style = (
        'UNIT - IV: E mail privacy: Pretty Good Privacy (PGP) and S/MIME. '
        'IP Security: Over view, IP Security Architecture, Authentication Header, '
        'Encapsulating Security Payload, Combining Security Association and Key Management.'
    )

    prompt = f"""You are an experienced university curriculum designer who writes official academic
course syllabi for engineering/degree programs (the kind used by Indian universities such as JNTU Hyderabad).

Generate a COMPLETE, DETAILED, and TECHNICALLY ACCURATE course syllabus in STRICT JSON for the course below.
Do not include any text, markdown, or code fences outside the JSON object. Return ONLY valid JSON.

COURSE DETAILS
- Course Name: {data.get('course_name')}
- Course Code: {data.get('course_code') or 'Not specified'}
- Level: {data.get('level')}
- Semester: {data.get('semester')}
- Credits: {data.get('credits')}
- Duration: {data.get('duration_weeks')} weeks
- Number of Units required: {unit_count}
- Instructor-provided rough objectives / focus areas: {data.get('course_objectives') or 'None provided - infer standard objectives for this course.'}
- Additional requirements from instructor: {data.get('additional_requirements') or 'None.'}

STRICT FORMATTING RULES
1. "prerequisites": a short paragraph (1-3 sentences) naming prerequisite courses/topics appropriate for the level.
2. "overview": a single flowing paragraph, no bullet points.
3. "objectives": 5-8 concise bullet points, each a complete sentence, written from the instructor's point of view
   (e.g. "To introduce students to ...", "To familiarize students with ...").
4. "outcomes": one bullet per Course Outcome, numbered "CO1:", "CO2:", etc., each starting with a measurable Bloom's
   taxonomy verb (Understand, Apply, Analyze, Design, Evaluate, Implement, Demonstrate...), written as something a
   STUDENT will be able to do after completing the course.
5. "units": exactly {unit_count} units labelled with roman numerals {unit_numerals}. Each unit's "topics" field MUST
   follow this exact academic style (sub-topic heading, colon, comma separated sub points, period, next sub-topic
   heading...) — for example:
   "{example_unit_style}"
   Cover the subject matter progressively from foundational (Unit I) to advanced (last unit), and make sure the
   total content is realistic for a {data.get('duration_weeks')}-week, {data.get('credits')}-credit course.
   Do NOT abbreviate, shorten, truncate, or drop sub-topics to save space, even with {unit_count} units to
   cover — every unit must be fully detailed. It is fine for the response to be long.
6. "assignments": exactly 5 short, practical, unit-specific assignment tasks per unit (things a student can actually
   do - implement, design, compare, report, demonstrate).
7. "unit_tests": exactly 2 unit-test exam questions per unit, written in formal exam language (e.g. "Explain ... with
   a suitable example.", "Differentiate between ... and ... .", "Design/Illustrate ...").
8. "text_books": 2-3 real-style textbook citations (Author, Title, Edition, Publisher, Year) relevant to this exact
   subject.
9. "reference_books": 2-3 additional real-style reference citations.
10. "electronic_resources": 2-4 realistic online resources (NPTEL course name, official documentation, well-known
    MOOC platform course, or standards body) relevant to this subject, each with a one-line description.

Return JSON matching exactly this shape (keys and types), with real generated content (not placeholders):
{schema}
"""
    return prompt


def extract_json(text):
    """Pull the first balanced {...} JSON object out of an LLM response."""
    text = text.strip()
    text = re.sub(r"^```(json)?", "", text.strip(), flags=re.IGNORECASE).strip()
    text = re.sub(r"```$", "", text.strip()).strip()
    start = text.find("{")
    if start == -1:
        raise ValueError("No JSON object found in model output")
    depth = 0
    for i in range(start, len(text)):
        if text[i] == "{":
            depth += 1
        elif text[i] == "}":
            depth -= 1
            if depth == 0:
                return json.loads(text[start : i + 1])
    raise ValueError("Unbalanced JSON in model output")


def call_llm(prompt):
    if not LLM_API_KEY:
        raise RuntimeError("LLM_API_KEY not configured")

    url = f"{LLM_BASE_URL}/chat/completions"
    headers = {
        "Authorization": f"Bearer {LLM_API_KEY}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": LLM_MODEL,
        "messages": [
            {
                "role": "system",
                "content": "You are a precise academic syllabus-writing assistant. You always reply with valid JSON only.",
            },
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.4,
        "max_tokens": 8000,
    }

    resp = requests.post(url, headers=headers, json=payload, timeout=LLM_TIMEOUT)
    resp.raise_for_status()
    body = resp.json()
    content = body["choices"][0]["message"]["content"]
    return extract_json(content)


# ---------------------------------------------------------------------------
# Deterministic fallback generator (used if no API key is set, or the LLM
# call / JSON parsing fails for any reason). This guarantees the app NEVER
# shows an error screen to the user.
# ---------------------------------------------------------------------------
def fallback_generate(data):
    course = data.get("course_name") or "the course"
    unit_count = int(data.get("unit_count") or DEFAULT_UNIT_COUNT)
    numerals = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"][:unit_count]
    level = data.get("level") or "Undergraduate"

    objectives_seed = data.get("course_objectives") or ""
    extra_obj = (
        [s.strip() for s in re.split(r"[\n.;]+", objectives_seed) if s.strip()]
        if objectives_seed
        else []
    )

    units = []
    for idx, num in enumerate(numerals, start=1):
        units.append(
            {
                "number": num,
                "title": f"Module {idx}: Foundations and Applications of {course} — Part {idx}",
                "topics": (
                    f"Introduction: Overview of key concept {idx} in {course}, motivation, "
                    f"real-world relevance, and historical context. "
                    f"Core Techniques: Fundamental methods and models used in this part of {course}, "
                    f"worked examples, and common pitfalls. "
                    f"Applications: Case studies and practical applications of unit {idx} concepts, "
                    f"tools used in industry, and a brief look ahead to the next unit."
                ),
                "assignments": [
                    f"Summarise the core concepts of Unit {num} in a short technical report.",
                    f"Implement/demonstrate one technique from Unit {num} using a tool of your choice.",
                    f"Compare two approaches discussed in Unit {num} and justify which is better for a given scenario.",
                    f"Prepare a presentation covering the real-world applications of Unit {num} topics.",
                    f"Solve a case-study problem based on Unit {num} and document your approach.",
                ],
                "unit_tests": [
                    f"Explain the key concepts covered in Unit {num} with suitable examples.",
                    f"Differentiate between any two major techniques discussed in Unit {num}.",
                ],
            }
        )

    objectives = extra_obj[:8] if extra_obj else []
    if len(objectives) < 5:
        objectives += [
            f"To introduce students to the fundamental concepts of {course}.",
            f"To build a strong theoretical foundation required for {course}.",
            f"To enable students to apply {course} concepts to solve practical problems.",
            f"To familiarise students with current tools and techniques used in {course}.",
            f"To develop analytical and design skills relevant to {course}.",
        ][: 5 - len(objectives)] if len(objectives) < 5 else []

    return {
        "prerequisites": (
            f"Basic knowledge of mathematics, fundamentals of computer science, and introductory "
            f"programming concepts is recommended before taking {course}, along with any "
            f"foundational subjects typically completed prior to this {level.lower()} course."
        ),
        "overview": (
            f"This course on {course} is designed to provide students with a comprehensive "
            f"understanding of both the theoretical foundations and practical applications of the "
            f"subject. Over the duration of the course, students will progress from fundamental "
            f"concepts to more advanced topics, gaining hands-on experience through assignments "
            f"and unit tests along the way. The course emphasises problem solving, critical "
            f"thinking, and the practical application of concepts to real-world scenarios. By the "
            f"end of the course, students will have developed the knowledge and skills needed to "
            f"confidently apply {course} concepts in academic projects and professional settings. "
            f"The syllabus has been structured to align with current industry practices and "
            f"academic standards."
        ),
        "objectives": objectives,
        "outcomes": [
            f"CO1: Understand the fundamental concepts and terminology used in {course}.",
            f"CO2: Apply core techniques of {course} to solve well-defined problems.",
            f"CO3: Analyse different approaches within {course} and select suitable methods for a given scenario.",
            f"CO4: Design simple solutions/systems based on the principles learnt in {course}.",
            f"CO5: Evaluate and demonstrate the practical applications of {course} through assignments and case studies.",
        ],
        "units": units,
        "text_books": [
            f'Author Name(s), "{course}: Principles and Practice", 3rd Edition, Reputed Publisher, 2022.',
            f'Author Name(s), "Introduction to {course}", 2nd Edition, Reputed Publisher, 2021.',
        ],
        "reference_books": [
            f'Author Name(s), "Advanced {course}", 1st Edition, Reputed Publisher, 2023.',
            f'Author Name(s), "{course} — A Practical Approach", 2nd Edition, Reputed Publisher, 2020.',
        ],
        "electronic_resources": [
            f"NPTEL course related to {course} — nptel.ac.in (video lectures and assignments).",
            f"Official documentation / standard reference material relevant to {course}.",
            f"Relevant MOOC (Coursera/edX) course covering {course} fundamentals and applications.",
        ],
    }


def generate_syllabus(data):
    """Try the LLM; fall back to the deterministic generator on ANY failure."""
    try:
        prompt = build_prompt(data)
        result = call_llm(prompt)
        result = normalize_syllabus(result, data)
        return result, True, None
    except Exception as exc:  # noqa: BLE001 - we want to catch everything here
        fallback = fallback_generate(data)
        fallback = normalize_syllabus(fallback, data)
        return fallback, False, str(exc)


def normalize_syllabus(result, data):
    """Validate and normalize AI/fallback syllabus output before displaying it."""

    unit_count = int(data.get("unit_count") or DEFAULT_UNIT_COUNT)
    numerals = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"][:unit_count]

    # Use the existing fallback generator as a source of safe defaults.
    baseline = fallback_generate(data)

    if not isinstance(result, dict):
        result = {}

    def safe_text(value, default=""):
        if isinstance(value, str) and value.strip():
            return value.strip()

        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return str(value)

        return default

    def safe_list(value, defaults):
        if not isinstance(value, list):
            value = []

        cleaned = [
            str(item).strip()
            for item in value
            if isinstance(item, (str, int, float))
            and not isinstance(item, bool)
            and str(item).strip()
        ]

        return cleaned if cleaned else list(defaults)

    # Validate main syllabus sections.
    for key in [
        "prerequisites",
        "overview",
    ]:
        result[key] = safe_text(
            result.get(key),
            baseline[key]
        )

    for key in [
        "objectives",
        "outcomes",
        "text_books",
        "reference_books",
        "electronic_resources",
    ]:
        result[key] = safe_list(
            result.get(key),
            baseline[key]
        )

    # Validate units.
    units = result.get("units")

    if not isinstance(units, list):
        units = []

    fixed_units = []

    for i, num in enumerate(numerals):
        if i < len(units) and isinstance(units[i], dict):
            unit = units[i]
        else:
            unit = {}

        default_unit = baseline["units"][i]

        assignments = safe_list(
            unit.get("assignments"),
            default_unit["assignments"]
        )

        unit_tests = safe_list(
            unit.get("unit_tests"),
            default_unit["unit_tests"]
        )

        # Ensure exactly 5 assignments.
        while len(assignments) < 5:
            assignments.append(
                default_unit["assignments"][len(assignments)]
            )

        assignments = assignments[:5]

        # Ensure exactly 2 unit-test questions.
        while len(unit_tests) < 2:
            unit_tests.append(
                default_unit["unit_tests"][len(unit_tests)]
            )

        unit_tests = unit_tests[:2]

        fixed_units.append({
            "number": safe_text(
                unit.get("number"),
                num
            ),
            "title": safe_text(
                unit.get("title"),
                default_unit["title"]
            ),
            "topics": safe_text(
                unit.get("topics"),
                default_unit["topics"]
            ),
            "assignments": assignments,
            "unit_tests": unit_tests,
        })

    result["units"] = fixed_units

    return result

def validate_generation_data(data):
    """Validate and normalize user input before syllabus generation."""

    if not isinstance(data, dict):
        return None, "Invalid request. Please submit valid course details."

    text_limits = {
        "course_name": 150,
        "course_code": 40,
        "level": 30,
        "semester": 50,
        "course_objectives": 4000,
        "additional_requirements": 4000,
    }

    cleaned = {}

    for field, max_length in text_limits.items():
        value = data.get(field, "")

        if value is None:
            value = ""

        if not isinstance(value, str):
            return None, f"Invalid value provided for {field.replace('_', ' ')}."

        value = value.strip()

        if len(value) > max_length:
            return None, (
                f"{field.replace('_', ' ').capitalize()} "
                f"cannot exceed {max_length} characters."
            )

        cleaned[field] = value

    # Required text fields
    required_fields = ["course_name", "level", "semester"]

    for field in required_fields:
        if not cleaned[field]:
            return None, f"{field.replace('_', ' ').capitalize()} is required."

    # Validate course level
    allowed_levels = ["Undergraduate", "Postgraduate", "Diploma"]

    if cleaned["level"] not in allowed_levels:
        return None, "Please select a valid course level."

    # Validate integer fields
    integer_fields = {
        "credits": (1, 10),
        "duration_weeks": (1, 52),
    }

    for field, (minimum, maximum) in integer_fields.items():
        value = data.get(field)

        if isinstance(value, bool):
            return None, f"{field.replace('_', ' ').capitalize()} must be a whole number."

        try:
            number = int(str(value).strip())
        except (TypeError, ValueError):
            return None, f"{field.replace('_', ' ').capitalize()} must be a whole number."

        if str(value).strip() != str(number):
            return None, f"{field.replace('_', ' ').capitalize()} must be a whole number."

        if not minimum <= number <= maximum:
            return None, (
                f"{field.replace('_', ' ').capitalize()} "
                f"must be between {minimum} and {maximum}."
            )

        cleaned[field] = number

    # Validate number of units
    unit_count = data.get("unit_count", DEFAULT_UNIT_COUNT)

    if unit_count is None or str(unit_count).strip() == "":
        unit_count = DEFAULT_UNIT_COUNT

    if isinstance(unit_count, bool):
        return None, "Please select a valid number of units."

    try:
        unit_count = int(str(unit_count).strip())
    except (TypeError, ValueError):
        return None, "Please select a valid number of units."

    if unit_count not in [4, 5, 6, 8]:
        return None, "Number of units must be 4, 5, 6, or 8."

    cleaned["unit_count"] = unit_count

    return cleaned, None
# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/generate", methods=["POST"])
def api_generate():
    data = request.get_json(silent=True)

    data, validation_error = validate_generation_data(data)

    if validation_error:
        return jsonify({
            "ok": False,
            "error": validation_error
        }), 400

    syllabus, ai_used, warning = generate_syllabus(data)
    meta = {
        "course_name": data.get("course_name"),
        "course_code": data.get("course_code") or "",
        "level": data.get("level"),
        "semester": data.get("semester"),
        "credits": data.get("credits"),
        "duration_weeks": data.get("duration_weeks"),
    }
    return jsonify(
        {
            "ok": True,
            "ai_generated": ai_used,
            "warning": None if ai_used else "AI service unavailable — a baseline syllabus was generated instead. You can edit everything below.",
            "meta": meta,
            "syllabus": syllabus,
        }
    )


@app.route("/api/export-docx", methods=["POST"])
def api_export_docx():
    if Document is None:
        return jsonify({"ok": False, "error": "python-docx not available on server"}), 500

    payload = request.get_json(force=True, silent=True) or {}
    meta = payload.get("meta") or {}
    syllabus = payload.get("syllabus") or {}

    try:
        buf = build_docx(meta, syllabus)
    except Exception as exc:  # noqa: BLE001
        return jsonify({"ok": False, "error": f"Could not build document: {exc}"}), 500

    filename = re.sub(r"[^A-Za-z0-9_-]+", "_", meta.get("course_name") or "Syllabus")
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"{filename}_Syllabus.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def build_docx(meta, syllabus):
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), "Times New Roman")

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    def heading(text, size=14):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(size)
        r.font.name = "Times New Roman"
        p.space_after = Pt(6)
        return p

    def para(text, justify=True):
        p = doc.add_paragraph(text)
        if justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    def bullet(text):
        doc.add_paragraph(text, style="List Bullet")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("COURSE SYLLABUS")
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = "Times New Roman"

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(meta.get("course_name", ""))
    r.bold = True
    r.font.size = Pt(15)

    info = doc.add_table(rows=0, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("Course Code", meta.get("course_code") or "—"),
        ("Level", meta.get("level") or "—"),
        ("Semester", meta.get("semester") or "—"),
        ("Credits", meta.get("credits") or "—"),
        ("Duration", f"{meta.get('duration_weeks') or '—'} weeks"),
    ]
    for k, v in rows:
        row = info.add_row()
        row.cells[0].text = k
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[1].text = str(v)

    doc.add_paragraph()

    heading("Prerequisites")
    para(syllabus.get("prerequisites", ""))

    heading("Course Overview")
    para(syllabus.get("overview", ""))

    heading("Course Objectives")
    for o in syllabus.get("objectives", []):
        bullet(o)

    heading("Course Outcomes")
    para("After completion of the course, students should be able to:", justify=False)
    for o in syllabus.get("outcomes", []):
        bullet(o)

    for unit in syllabus.get("units", []):
        doc.add_paragraph()
        heading(f"UNIT - {unit.get('number', '')}: {unit.get('title', '')}")
        para(unit.get("topics", ""))

        sub_h = doc.add_paragraph()
        r = sub_h.add_run("Assignments")
        r.bold = True
        r.underline = True
        for i, a in enumerate(unit.get("assignments", []), start=1):
            bullet(f"{i}. {a}")

        sub_h2 = doc.add_paragraph()
        r = sub_h2.add_run("Unit Test Questions")
        r.bold = True
        r.underline = True
        for i, q in enumerate(unit.get("unit_tests", []), start=1):
            bullet(f"{i}. {q}")

    doc.add_paragraph()
    heading("Text Books")
    for i, b in enumerate(syllabus.get("text_books", []), start=1):
        para(f"{i}. {b}", justify=False)

    heading("Reference Books")
    for i, b in enumerate(syllabus.get("reference_books", []), start=1):
        para(f"{i}. {b}", justify=False)

    heading("Electronic Resources")
    for i, b in enumerate(syllabus.get("electronic_resources", []), start=1):
        para(f"{i}. {b}", justify=False)

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(f"Generated on {datetime.date.today().strftime('%d %B %Y')}")
    r.italic = True
    r.font.size = Pt(9)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


@app.route("/healthz")
def healthz():
    return jsonify({"status": "ok", "llm_configured": bool(LLM_API_KEY)})


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=bool(os.environ.get("FLASK_DEBUG")))
